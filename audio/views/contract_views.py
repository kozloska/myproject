import io
import re
import zipfile
from copy import deepcopy
from datetime import datetime, date

from django.http import JsonResponse, HttpResponse
from django.views import View
from django.views.decorators.csrf import csrf_exempt
from django.utils.decorators import method_decorator

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

# ===========================================================================
# Сумма прописью
# ===========================================================================

_ONES_MALE = ['', 'один', 'два', 'три', 'четыре', 'пять', 'шесть', 'семь', 'восемь', 'девять']
_ONES_FEMALE = ['', 'одна', 'две', 'три', 'четыре', 'пять', 'шесть', 'семь', 'восемь', 'девять']
_TEENS = ['десять', 'одиннадцать', 'двенадцать', 'тринадцать', 'четырнадцать',
          'пятнадцать', 'шестнадцать', 'семнадцать', 'восемнадцать', 'девятнадцать']
_TENS = ['', '', 'двадцать', 'тридцать', 'сорок', 'пятьдесят',
         'шестьдесят', 'семьдесят', 'восемьдесят', 'девяносто']
_HUNDREDS = ['', 'сто', 'двести', 'триста', 'четыреста', 'пятьсот',
             'шестьсот', 'семьсот', 'восемьсот', 'девятьсот']

_GROUPS = [
    (10 ** 9, ('миллиард', 'миллиарда', 'миллиардов'), False),
    (10 ** 6, ('миллион', 'миллиона', 'миллионов'), False),
    (10 ** 3, ('тысяча', 'тысячи', 'тысяч'), True),
]


def _plural(n, forms):
    n = abs(int(n)) % 100
    if 11 <= n <= 14:
        return forms[2]
    n2 = n % 10
    if n2 == 1:
        return forms[0]
    if 2 <= n2 <= 4:
        return forms[1]
    return forms[2]


def _three_digits_to_words(n, feminine=False):
    words = []
    h, r = divmod(n, 100)
    if h:
        words.append(_HUNDREDS[h])
    if 10 <= r < 20:
        words.append(_TEENS[r - 10])
    else:
        t, o = divmod(r, 10)
        if t:
            words.append(_TENS[t])
        if o:
            words.append((_ONES_FEMALE if feminine else _ONES_MALE)[o])
    return words


def num_to_words_ru(n):
    n = int(n)
    if n == 0:
        return 'ноль'
    words = []
    remaining = n
    for div, forms, feminine in _GROUPS:
        if remaining >= div:
            count = remaining // div
            remaining %= div
            words.extend(_three_digits_to_words(count, feminine=feminine))
            words.append(_plural(count, forms))
    if remaining:
        words.extend(_three_digits_to_words(remaining, feminine=False))
    return ' '.join(words)


def sum_to_full_str(amount):
    amount = float(amount)
    rub = int(amount)
    kop = int(round((amount - rub) * 100))

    if amount == rub:
        amount_display = str(rub)
    else:
        amount_display = f"{amount:.2f}".rstrip('0').rstrip('.')

    words = num_to_words_ru(rub)
    rub_word = _plural(rub, ('рубль', 'рубля', 'рублей'))
    kop_word = _plural(kop, ('копейка', 'копейки', 'копеек'))

    return f"{amount_display} ({words} {rub_word} {kop:02d} {kop_word})"


# ===========================================================================
# Мелкие помощники
# ===========================================================================

class ContractError(Exception):
    """Ошибка пользовательских данных -> ответ 400 с текстом."""


def _clean(value):
    """Убирает лишний .0 у чисел, приводит к строке, режет пробелы."""
    if value is None:
        return ''
    if isinstance(value, float):
        if value == int(value):
            return str(int(value))
        return str(value)
    return str(value).strip()


def _format_passport_series(value):
    """2515 -> '25 15' (стандартный вид серии паспорта РФ)."""
    s = _clean(value)
    if len(s) == 4 and s.isdigit():
        return f"{s[:2]} {s[2:]}"
    return s


def _fmt_date(value):
    if isinstance(value, (datetime, date)):
        return value.strftime('%d.%m.%Y')
    return _clean(value)


def _fmt_contract_number(value):
    """7 -> '07', 12 -> '12' (как в номерах ГПХ: ЦК-ПИИИ-07/26)."""
    try:
        return f"{int(float(value)):02d}"
    except (TypeError, ValueError):
        return _clean(value)


def _to_amount(value):
    if isinstance(value, str):
        value = value.replace('\xa0', '').replace(' ', '').replace(',', '.')
    return float(value)


def make_initials(fio):
    """'Козлова Вероника Александровна' -> 'В.А. Козлова'"""
    parts = [p for p in fio.split() if p]
    if not parts:
        return ''
    surname = parts[0]
    initials = ''.join(f"{p[0].upper()}." for p in parts[1:3])
    return f"{initials} {surname}".strip() if initials else surname


def _safe_filename(s):
    return re.sub(r'[\\/*?:"<>|\s]+', '_', s).strip('_')


# ===========================================================================
# Замена плейсхолдеров в docx
# ===========================================================================
# Понимает {name} в любом регистре и с пробелами внутри скобок ({ initials }),
# а также "голые" sum / fio (в текущем шаблоне ГПХ они написаны без скобок).
# Плейсхолдер, которого нет в mapping, остаётся в тексте как есть.
# Ключи mapping — в нижнем регистре.

_TOKEN_RE = re.compile(r'\{\s*([^{}\s]+)\s*\}|(?<![\w{}])(sum|fio)(?![\w{}])', re.IGNORECASE)


def _replace_in_paragraph(paragraph, mapping):
    runs = paragraph.runs
    if not runs:
        return
    texts = [r.text for r in runs]
    full = ''.join(texts)

    edits = []
    for m in _TOKEN_RE.finditer(full):
        key = (m.group(1) or m.group(2)).lower()
        if key in mapping:
            edits.append((m.start(), m.end(), str(mapping[key])))
    if not edits:
        return

    spans, pos = [], 0
    for t in texts:
        spans.append((pos, pos + len(t)))
        pos += len(t)

    new = list(texts)
    # с конца к началу, чтобы смещения предыдущих правок не ломались
    for start, end, value in reversed(edits):
        covered = [i for i, (s, e) in enumerate(spans) if s < end and e > start]
        first, last = covered[0], covered[-1]
        if first == last:
            off = spans[first][0]
            new[first] = new[first][:start - off] + value + new[first][end - off:]
        else:
            new[first] = new[first][:start - spans[first][0]] + value
            new[last] = new[last][end - spans[last][0]:]
            for i in range(first + 1, last):
                new[i] = ''

    for i, run in enumerate(runs):
        if new[i] != texts[i]:
            run.text = new[i]


def _set_paragraph_text(paragraph, text):
    """Меняет текст абзаца, сохраняя форматирование первого run."""
    if paragraph.runs:
        paragraph.runs[0].text = text
        for r in paragraph.runs[1:]:
            r.text = ''
    else:
        paragraph.add_run(text)


def _iter_paragraphs(container, _seen=None):
    """Все абзацы документа/ячейки, включая вложенные таблицы."""
    seen = _seen if _seen is not None else set()
    for p in container.paragraphs:
        yield p
    for table in container.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell._tc in seen:  # объединённые ячейки
                    continue
                seen.add(cell._tc)
                yield from _iter_paragraphs(cell, seen)


def _iter_all_tables(container):
    for table in container.tables:
        yield table
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_all_tables(cell)


# ===========================================================================
# ГПХ: размножение строк таблиц под несколько групп
# ===========================================================================

_GROUP_RE = re.compile(r'\{\s*group\s*\}', re.IGNORECASE)
_DISC_RE = re.compile(r'\{\s*discipline\s*\}', re.IGNORECASE)
_HOURS_COST_RE = re.compile(r'\{\s*(hours|cost)\s*\}', re.IGNORECASE)
_STAGE_NUM_RE = re.compile(r'^\s*(I|\d+)\s*\.?\s*$')


def _tr_text(tr):
    return ''.join(t.text or '' for t in tr.iter(qn('w:t')))


def _fill_tr(tr, table, mapping):
    for p in tr.iter(qn('w:p')):
        _replace_in_paragraph(Paragraph(p, table), mapping)


def _renumber_stages(table):
    """
    Перенумеровывает первый столбец («I.», «2.», «3.» …) после размножения строк.
    Первая строка сохраняет подпись из шаблона (там «I.»).
    """
    numbered = []
    for tr in table._tbl.tr_lst:
        tcs = tr.tc_lst
        if not tcs:
            continue
        cell_text = ''.join(t.text or '' for t in tcs[0].iter(qn('w:t')))
        if _STAGE_NUM_RE.match(cell_text):
            numbered.append((tcs[0], cell_text))

    for k, (tc, original) in enumerate(numbered):
        label = original.strip() if k == 0 else f"{k + 1}."
        if tc.p_lst and ''.join(t.text or '' for t in tc.iter(qn('w:t'))).strip() != label:
            _set_paragraph_text(Paragraph(tc.p_lst[0], table), label)


def _money(value):
    value = round(float(value), 2)
    if value == int(value):
        return f"{int(value):,}".replace(',', ' ')
    return f"{value:,.2f}".replace(',', ' ')


def _hours_text(hours):
    hours = float(hours)
    if hours == int(hours):
        n = int(hours)
        return f"{n} {_plural(n, ('час', 'часа', 'часов'))}"
    return f"{str(hours).replace('.', ',')} часа"


def _row_instances(tr, group, program):
    """
    Какие строки получатся из строки-прототипа для одной группы:
      • есть {discipline}        -> по строке на каждую дисциплину программы
      • есть {hours}/{cost}      -> по строке на каждый этап «сопровождение»
      • только {group} (или без) -> одна строка
    Возвращает список словарей-подстановок.
    """
    text = _tr_text(tr)
    if program and _DISC_RE.search(text):
        return [{'group': group, 'discipline': d['name'], 'hours': _hours_text(d['hours']),
                 'cost': _money(d['hours'] * d['rate'])} for d in program['lessons']]
    if program and _HOURS_COST_RE.search(text):
        return [{'group': group, 'hours': _hours_text(d['hours']),
                 'cost': _money(d['hours'] * d['rate'])} for d in program['support']]
    return [{'group': group}]


def _expand_group_rows(table, groups, program=None):
    """
    Каждая подряд идущая пачка строк, где встречается {group}/{discipline}/{hours}/{cost},
    пересобирается для каждой группы:
      [занятия(по дисциплине), сопровождение]  ->  гр.A: дисциплина 1, дисциплина 2, …, сопровождение;
                                                   гр.B: дисциплина 1, дисциплина 2, …, сопровождение
    Тексты, сроки и оформление берутся из строк-прототипов шаблона.
    """
    def is_stage_row(tr):
        t = _tr_text(tr)
        return bool(_GROUP_RE.search(t) or _DISC_RE.search(t) or _HOURS_COST_RE.search(t))

    blocks, cur = [], []
    for tr in list(table._tbl.tr_lst):
        if is_stage_row(tr):
            cur.append(tr)
        elif cur:
            blocks.append(cur)
            cur = []
    if cur:
        blocks.append(cur)
    if not blocks:
        return

    for block in blocks:
        new_rows = []
        for g in groups:
            for tr in block:
                for mapping in _row_instances(tr, g, program):
                    new_tr = deepcopy(tr)
                    _fill_tr(new_tr, table, mapping)
                    new_rows.append(new_tr)
        anchor = block[-1]
        for new_tr in reversed(new_rows):
            anchor.addnext(new_tr)
        for tr in block:
            tr.getparent().remove(tr)

    _renumber_stages(table)


def _fill_docx(template_bytes, mapping, groups=None, program=None):
    doc = Document(io.BytesIO(template_bytes))

    if groups and (len(groups) > 1 or program):
        for table in _iter_all_tables(doc):
            _expand_group_rows(table, groups, program)

    containers = [doc] + [s.header for s in doc.sections] + [s.footer for s in doc.sections]
    for c in containers:
        for p in _iter_paragraphs(c):
            _replace_in_paragraph(p, mapping)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ===========================================================================
# Генераторы по типам договора
# ===========================================================================

def _generate_fiz(excel_bytes, template_bytes, start_number, **_):
    """
    Договор с физлицом (платные образовательные услуги).
    Excel: лист «Лист1», данные с 3-й строки, колонки фиксированные:
    A ФИО | B Сумма | C Скидка % | D Серия | E Номер | F Кем выдан | G Код |
    H СНИЛС | I ИНН | J Телефон | K Email
    """
    SHEET = 'Лист1'
    FIO, SUM, DISCOUNT, P_SER, P_NUM, ISSUED, CODE, SNILS, INN, TEL, EMAIL = range(11)

    wb = load_workbook(io.BytesIO(excel_bytes), data_only=True)
    if SHEET not in wb.sheetnames:
        raise ContractError(f"Лист '{SHEET}' не найден. Доступные: {wb.sheetnames}")

    year = datetime.now().year
    number = start_number
    result = []

    for row in wb[SHEET].iter_rows(min_row=3, values_only=True):
        if len(row) <= EMAIL:
            row = tuple(row) + (None,) * (EMAIL + 1 - len(row))

        fio = _clean(row[FIO])
        if not fio or row[SUM] is None:
            continue

        # Столбец C («Скидка») не используется и в договор не попадает.
        sum_str = sum_to_full_str(_to_amount(row[SUM]))

        mapping = {
            'number': str(number),
            'year': str(year),
            'fio': fio,
            'initials': make_initials(fio),
            'sum': sum_str,
            'passport': f"Серия: {_format_passport_series(row[P_SER])} Номер: {_clean(row[P_NUM])}",
            'lssued': _clean(row[ISSUED]),
            'code': _clean(row[CODE]),
            'snils': _clean(row[SNILS]),
            'inn': _clean(row[INN]),
            'telephone': _clean(row[TEL]),
            'email': _clean(row[EMAIL]),
        }

        filename = f"Договор_{number}_{_safe_filename(fio)}.docx"
        result.append((filename, _fill_docx(template_bytes, mapping)))
        number += 1

    if not result:
        raise ContractError('Не найдено данных для генерации. Проверьте, что в листе есть ФИО и стоимость.')
    return result


# --- ГПХ -------------------------------------------------------------------

def _norm_header(v):
    return re.sub(r'\s+', ' ', _clean(v).lower().replace('ё', 'е'))


# поле -> допустимые названия заголовков в Excel (регистр/ё не важны)
_GPH_COLUMNS = {
    'fio': ['фио', 'ф.и.о.', 'фамилия имя отчество'],
    'group': ['группа', 'группы', 'группа(ы)', 'группа (группы)'],
    'sum': ['сумма договора', 'сумма', 'стоимость'],
    'birthday': ['дата рождения'],
    'passport': ['номер и серия паспорта', 'серия и номер паспорта', 'паспорт'],
    'lssued': ['кем выдан паспорт', 'кем выдан'],
    'datepassport': ['дата выдачи паспорта', 'дата выдачи'],
    'snils': ['снилс'],
    'inn': ['инн'],
    'address': ['адрес регистрации', 'адрес'],
    'work': ['место работы', 'место работы, должность'],
    'bank': ['банк', 'банк-получатель'],
    'bik': ['бик'],
    'kpp': ['кпп'],
    'ks': ['к/с', 'кс', 'корр. счет', 'корреспондентский счет'],
    'rs': ['р/с', 'рс', 'расчетный счет'],
    'telephone': ['номер телефона', 'телефон'],
    'number': ['цифра договора', 'номер договора'],
    'direction': ['программа', 'направление', 'направление подготовки', 'название программы'],
    'letter': ['буква', 'литера', 'шифр'],
}
_GPH_ALIAS = {alias: field for field, aliases in _GPH_COLUMNS.items() for alias in aliases}


def _find_header(ws, alias_map, required, max_rows=10):
    """Ищет строку заголовков; возвращает (номер_строки, {индекс_столбца: поле}) или None."""
    for header_no, header in enumerate(ws.iter_rows(min_row=1, max_row=max_rows, values_only=True), start=1):
        col_map = {}
        for idx, cell in enumerate(header):
            field = alias_map.get(_norm_header(cell))
            if field and field not in col_map.values():
                col_map[idx] = field
        if required in col_map.values():
            return header_no, col_map
    return None


def _read_table(ws, alias_map, required):
    found = _find_header(ws, alias_map, required)
    if not found:
        return None
    header_no, col_map = found
    rows = []
    for row_no, row in enumerate(ws.iter_rows(min_row=header_no + 1, values_only=True), start=header_no + 1):
        rows.append((row_no, {f: row[i] for i, f in col_map.items() if i < len(row)}))
    return rows


PROGRAMS_SHEET = 'программы'

_PROGRAM_COLUMNS = {
    'program': ['программа', 'направление', 'название программы'],
    'letter': ['буква', 'литера', 'шифр'],
    'kind': ['тип этапа', 'тип', 'вид'],
    'discipline': ['дисциплина', 'название дисциплины', 'модуль'],
    'hours': ['часы', 'часов', 'количество часов'],
    'rate': ['ставка, руб./час', 'ставка', 'ставка руб/час', 'ставка, руб/час'],
}
_PROGRAM_ALIAS = {a: f for f, aliases in _PROGRAM_COLUMNS.items() for a in aliases}

DEFAULT_RATES = {'lesson': 900, 'support': 499}  # руб./час, если ставка не указана


def _read_programs(wb):
    """
    Лист «Программы»: одна строка = один этап (дисциплина или сопровождение).
    Возвращает {нормализованное_название: {name, letter, lessons[], support[]}} или None, если листа нет.
    """
    ws = next((w for w in wb.worksheets if _norm_header(w.title) == PROGRAMS_SHEET), None)
    if ws is None:
        return None
    rows = _read_table(ws, _PROGRAM_ALIAS, 'program')
    if rows is None:
        raise ContractError('На листе «Программы» не найден столбец «Программа».')

    programs = {}
    for row_no, d in rows:
        name = _clean(d.get('program'))
        if not name:
            continue
        p = programs.setdefault(_norm_header(name), {'name': name, 'letter': '', 'lessons': [], 'support': []})
        if not p['letter']:
            p['letter'] = _clean(d.get('letter'))

        kind = 'support' if _norm_header(d.get('kind')).startswith('сопровожд') else 'lesson'
        try:
            hours = _to_amount(d.get('hours'))
        except (TypeError, ValueError):
            raise ContractError(f'Лист «Программы», строка {row_no}: не указаны часы.')
        rate_raw = d.get('rate')
        try:
            rate = _to_amount(rate_raw) if _clean(rate_raw) != '' else DEFAULT_RATES[kind]
        except ValueError:
            raise ContractError(f'Лист «Программы», строка {row_no}: ставка «{rate_raw}» — не число.')

        if kind == 'lesson':
            disc = _clean(d.get('discipline'))
            if not disc:
                raise ContractError(f'Лист «Программы», строка {row_no}: не указана дисциплина.')
            p['lessons'].append({'name': disc, 'hours': hours, 'rate': rate})
        else:
            p['support'].append({'hours': hours, 'rate': rate})

    for p in programs.values():
        if not p['lessons']:
            raise ContractError(f"Программа «{p['name']}»: нет ни одной дисциплины.")
    return programs


def _generate_gph(excel_bytes, template_bytes, start_number, direction='', letter='', **_):
    """
    Договор ГПХ.
    Excel:
      • лист со слушателями (любое название, кроме «Программы»): заголовки в первых 10 строках
        — ФИО, Группа, Программа (или Направление), Сумма договора, паспорт и т.д.
      • лист «Программы» (необязательный): дисциплины, часы и ставки для каждой программы —
        по нему строятся строки таблиц Технического задания.
    Если листа «Программы» нет, работает прежняя схема: таблицы шаблона размножаются
    только по группам, а направление/буква берутся из столбцов или из формы.
    """
    year = datetime.now().year
    auto_number = start_number
    result = []

    wb = load_workbook(io.BytesIO(excel_bytes), data_only=True)
    programs = _read_programs(wb)

    rows = None
    for ws in wb.worksheets:
        if _norm_header(ws.title) == PROGRAMS_SHEET:
            continue
        rows = _read_table(ws, _GPH_ALIAS, 'fio')
        if rows is not None:
            break
    if rows is None:
        raise ContractError('Не найден столбец «ФИО» в заголовках Excel (ищу в первых 10 строках).')

    for row_no, data in rows:
        fio = _clean(data.get('fio'))
        if not fio:
            continue

        groups = [g.strip() for g in re.split(r'[,;\n]+', _clean(data.get('group'))) if g.strip()]
        if not groups:
            raise ContractError(f'Строка {row_no} ({fio}): не указана группа.')

        row_direction = _clean(data.get('direction')) or direction
        program = None
        if programs is not None:
            if not row_direction:
                raise ContractError(f'Строка {row_no} ({fio}): не указана программа.')
            program = programs.get(_norm_header(row_direction))
            if program is None:
                raise ContractError(f'Строка {row_no} ({fio}): программы «{row_direction}» нет на листе «Программы». '
                                    f'Доступные: {", ".join(p["name"] for p in programs.values())}')
            row_direction = program['name']

        row_letter = _clean(data.get('letter')) or (program['letter'] if program else '') or letter
        if not row_direction:
            raise ContractError(f'Строка {row_no} ({fio}): не указано направление '
                                f'(столбец «Направление» или поле в форме).')
        if not row_letter:
            raise ContractError(f'Строка {row_no} ({fio}): не указана буква для номера договора '
                                f'(столбец «Буква», лист «Программы» или поле в форме).')

        # сумма всегда берётся из Excel
        if _clean(data.get('sum')) == '':
            raise ContractError(f'Строка {row_no} ({fio}): не указана сумма договора.')
        try:
            sum_str = sum_to_full_str(_to_amount(data['sum']))
        except (TypeError, ValueError):
            raise ContractError(f'Строка {row_no} ({fio}): сумма «{data["sum"]}» — не число.')

        if _clean(data.get('number')) != '':
            number = _fmt_contract_number(data['number'])
        else:
            number = _fmt_contract_number(auto_number)
            auto_number += 1

        mapping = {
            'number': number,
            'letter': row_letter,
            'direction': row_direction,
            'year': str(year),
            'fio': fio,
            'initials': make_initials(fio),
            'group': ', '.join(groups),           # для абзацев; в таблицах — по строке на группу
            'modules': ', '.join(str(i + 1) for i in range(len(program['lessons']))) if program else '',
            'sum': sum_str,
            'birthday': _fmt_date(data.get('birthday')),
            'passport': _clean(data.get('passport')),
            'lssued': _clean(data.get('lssued')),
            'issued': _clean(data.get('lssued')),
            'datepassport': _fmt_date(data.get('datepassport')),
            'snils': _clean(data.get('snils')),
            'inn': _clean(data.get('inn')),
            'address': _clean(data.get('address')),
            'work': _clean(data.get('work')),
            'bank': _clean(data.get('bank')),
            'bik': _clean(data.get('bik')),
            'kpp': _clean(data.get('kpp')),
            'ks': _clean(data.get('ks')),
            'rs': _clean(data.get('rs')),
            'telephone': _clean(data.get('telephone')),
        }

        filename = f"Договор_ЦК-{_safe_filename(row_letter)}-{number}_26_{_safe_filename(fio)}.docx"
        result.append((filename, _fill_docx(template_bytes, mapping, groups=groups, program=program)))

    if not result:
        raise ContractError('Не найдено данных для генерации. Проверьте, что в листе есть ФИО.')
    return result


_GENERATORS = {
    'fiz': _generate_fiz,
    'gph': _generate_gph,
}


# ===========================================================================
# View
# ===========================================================================

@method_decorator(csrf_exempt, name='dispatch')
class ContractGenerate(View):
    """
    POST multipart/form-data:
      excel_file, template_file   — обязательно
      contract_type               — 'fiz' (по умолчанию) | 'gph'
      start_number                — номер первого договора (по умолчанию 1)
      direction, letter           — только для 'gph', запасной вариант, если нет столбцов в Excel
    """

    def post(self, request):
        try:
            excel_file = request.FILES.get('excel_file')
            template_file = request.FILES.get('template_file')
            if not excel_file or not template_file:
                return JsonResponse({'error': 'Необходимо загрузить Excel и шаблон Word'}, status=400)

            contract_type = request.POST.get('contract_type', 'fiz')
            generator = _GENERATORS.get(contract_type)
            if generator is None:
                return JsonResponse({'error': f"Неизвестный тип договора: '{contract_type}'"}, status=400)

            try:
                start_number = int(request.POST.get('start_number', 1))
            except (TypeError, ValueError):
                start_number = 1

            files = generator(
                excel_file.read(),
                template_file.read(),
                start_number,
                direction=request.POST.get('direction', '').strip(),
                letter=request.POST.get('letter', '').strip(),
            )

            zip_buffer = io.BytesIO()
            with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as zf:
                for name, content in files:
                    zf.writestr(name, content)

            response = HttpResponse(zip_buffer.getvalue(), content_type='application/zip')
            response['Content-Disposition'] = 'attachment; filename="contracts.zip"'
            return response

        except ContractError as e:
            return JsonResponse({'error': str(e)}, status=400)
        except Exception as e:
            import traceback
            traceback.print_exc()
            return JsonResponse({'error': f'Ошибка сервера: {str(e)}'}, status=500)


# ===========================================================================
# Образцы Excel для скачивания
# ===========================================================================

_X_FONT = 'Arial'
_X_BAND = PatternFill('solid', fgColor='2F5D6B')   # верхняя «шапка»-группа
_X_REQ = PatternFill('solid', fgColor='FCE4A8')    # обязательные столбцы
_X_OPT = PatternFill('solid', fgColor='E3ECF0')    # необязательные
_X_ZEBRA = PatternFill('solid', fgColor='F7FAFB')
_X_SIDE = Side(style='thin', color='B8C4CA')
_X_BORDER = Border(left=_X_SIDE, right=_X_SIDE, top=_X_SIDE, bottom=_X_SIDE)


def _x_fit(ws):
    """Альбомная ориентация и «по ширине страницы» — чтобы образец нормально печатался."""
    from openpyxl.worksheet.properties import PageSetupProperties
    ws.page_setup.orientation = 'landscape'
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0


def _x_band(ws, row, bands):
    """bands: [(первый_столбец, последний_столбец, текст)] — объединённая цветная шапка-группа."""
    for c1, c2, text in bands:
        ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
        for col in range(c1, c2 + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = _X_BAND
            cell.border = _X_BORDER
        cell = ws.cell(row=row, column=c1, value=text)
        cell.font = Font(name=_X_FONT, bold=True, color='FFFFFF', size=11)
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[row].height = 22


def _x_header(ws, row, items):
    """items: [(название, ширина, обязательный)]"""
    for i, (name, width, required) in enumerate(items, 1):
        c = ws.cell(row=row, column=i, value=name)
        c.font = Font(name=_X_FONT, bold=True, size=10, color='1F2D33')
        c.fill = _X_REQ if required else _X_OPT
        c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        c.border = _X_BORDER
        ws.column_dimensions[c.column_letter].width = width
    ws.row_dimensions[row].height = 34


def _x_row(ws, row, values, text_cols=(), zebra=False):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = Font(name=_X_FONT, size=10)
        c.border = _X_BORDER
        c.alignment = Alignment(wrap_text=True, vertical='center')
        if zebra:
            c.fill = _X_ZEBRA
        if i in text_cols:
            c.number_format = '@'
            c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='left')


def _x_notes(wb, title, lines):
    wi = wb.create_sheet('Памятка')
    wi.sheet_view.showGridLines = False
    wi.column_dimensions['A'].width = 4
    wi.column_dimensions['B'].width = 110
    wi.cell(row=1, column=2, value=title).font = Font(name=_X_FONT, bold=True, size=14, color='2F5D6B')
    wi.row_dimensions[1].height = 28
    for r, text in enumerate(lines, 3):
        c = wi.cell(row=r, column=2, value=text)
        c.font = Font(name=_X_FONT, size=10)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        wi.row_dimensions[r].height = 15 * (1 + len(text) // 105)


def _build_fiz_example():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Лист1'          # название листа важно: генератор ищет именно «Лист1»
    ws.sheet_view.showGridLines = False

    _x_band(ws, 1, [(1, 3, 'Договор'), (4, 7, 'Паспорт'), (8, 11, 'Документы и контакты')])
    _x_header(ws, 2, [
        ('ФИО полностью', 36, True), ('Стоимость обучения, руб.', 16, True), ('Скидка, % (не используется)', 16, False),
        ('Серия', 10, False), ('Номер', 12, False), ('Кем выдан', 36, False), ('Код подразд.', 13, False),
        ('СНИЛС', 17, False), ('ИНН', 15, False), ('Телефон', 18, False), ('Email', 26, False),
    ])
    people = [
        ['Орлов Дмитрий Сергеевич', 48760, None, '1122', '334455', 'ОУФМС России по г. Москве', '770-001',
         '112-233-445 95', '770100000001', '+7 900 111-22-33', 'orlov@example.com'],
        ['Смирнова Анна Павловна', 48760, 10, '2233', '445566', 'ГУ МВД России по г. Санкт-Петербургу', '780-002',
         '223-344-556 06', '780200000002', '+7 900 222-33-44', 'smirnova@example.com'],
    ]
    for r, row in enumerate(people, 3):
        _x_row(ws, r, row, text_cols=(4, 5, 7, 8, 9), zebra=(r % 2 == 0))
    _x_fit(ws)
    ws.freeze_panes = 'B3'
    ws['A2'].comment = Comment('Данные начинаются с 3-й строки. Лист должен называться «Лист1». '
                               'Строки-примеры замените своими.', 'Генератор')
    ws['B2'].comment = Comment('Стоимость обучения. Подставляется в договор числом и прописью.', 'Генератор')
    ws['C2'].comment = Comment('Столбец не используется и в договор не попадает. Оставьте пустым.', 'Генератор')

    _x_notes(wb, 'Как заполнять (договор с физическим лицом)', [
        '1. Одна строка = один слушатель = один договор. Данные начинаются с 3-й строки, лист называется «Лист1».',
        '2. Желтые заголовки — обязательные столбцы (ФИО, Стоимость обучения). Остальные можно оставлять пустыми.',
        '3. В «Стоимости обучения» указывайте просто сумму обучения. Столбец «Скидка» не используется — его можно оставить пустым.',
        '4. Серию паспорта пишите 4 цифрами без пробела (1122) — в договоре она будет выведена как «11 22».',
        '5. Строки-примеры замените своими данными. Порядок столбцов менять нельзя.',
    ])
    return wb


def _build_gph_example():
    wb = Workbook()

    # --- Слушатели
    ws = wb.active
    ws.title = 'Слушатели'
    ws.sheet_view.showGridLines = False
    _x_band(ws, 1, [(1, 4, 'Обязательно'), (5, 11, 'Личные данные и паспорт'),
                    (12, 17, 'Работа и банковские реквизиты'), (18, 19, 'Связь и номер')])
    _x_header(ws, 2, [
        ('ФИО', 34, True), ('Группа', 24, True), ('Программа', 44, True), ('Сумма договора', 16, True),
        ('Дата рождения', 14, False), ('Номер и серия паспорта', 20, False),
        ('Кем выдан паспорт', 34, False), ('Дата выдачи паспорта', 16, False),
        ('СНИЛС', 16, False), ('ИНН', 15, False), ('Адрес регистрации', 40, False),
        ('Место работы', 32, False), ('Банк', 22, False), ('БИК', 12, False), ('КПП', 12, False),
        ('К/с', 24, False), ('Р/с', 24, False), ('Номер телефона', 18, False), ('Цифра договора', 12, False),
    ])
    p1, p2 = 'Цифровые технологии в экономике', 'Основы программирования на Python'
    people = [
        ['Орлов Дмитрий Сергеевич', 'ЦТЭ-26-1', p1, 55960,
         '12.04.1988', '11 22 334455', 'ОУФМС России по г. Москве', '20.05.2008', '112-233-445 95',
         '770100000001', 'г. Москва, ул. Садовая, д. 5, кв. 12', 'Кафедра информатики', 'ПАО «Банк Пример»',
         '044525000', None, '30101810000000000000', '40817810000000000001', '+7 900 111-22-33', None],
        ['Смирнова Анна Павловна', 'ОПП-26-1', p2, 48760,
         '03.09.1992', '22 33 445566', 'ГУ МВД России по г. Санкт-Петербургу', '15.10.2012', '223-344-556 06',
         '780200000002', 'г. Санкт-Петербург, пр. Невский, д. 10, кв. 3', 'Кафедра математики', 'ПАО «Банк Пример»',
         '044525000', None, '30101810000000000000', '40817810000000000002', '+7 900 222-33-44', None],
        ['Волков Игорь Андреевич', 'ЦТЭ-26-1, ЦТЭ-26-2', p1, 111920,
         '27.11.1985', '33 44 556677', 'УВД г. Казани', '01.12.2005', '334-455-667 17',
         '165500000003', 'г. Казань, ул. Кремлёвская, д. 1, кв. 7', 'Кафедра экономики', 'ПАО «Банк Пример»',
         '044525000', None, '30101810000000000000', '40817810000000000003', '+7 900 333-44-55', None],
    ]
    for r, row in enumerate(people, 3):
        _x_row(ws, r, row, text_cols=(6, 9, 10, 14, 15, 16, 17, 18), zebra=(r % 2 == 0))
        ws.row_dimensions[r].height = 30
    _x_fit(ws)
    ws.freeze_panes = 'B3'
    ws['B2'].comment = Comment('Несколько групп — через запятую (как у третьего слушателя). '
                               'В таблицах ТЗ этапы повторятся для каждой группы.', 'Генератор')
    ws['C2'].comment = Comment('Должно совпадать с названием в столбце «Программа» на листе «Программы» (регистр не важен).', 'Генератор')
    ws['D2'].comment = Comment('Обязательно. Подставляется в договор и акт числом и прописью.', 'Генератор')
    ws['S2'].comment = Comment('Необязательно. Если пусто — сквозная нумерация с номера, указанного в форме.', 'Генератор')

    # --- Программы
    wp = wb.create_sheet('Программы')
    wp.sheet_view.showGridLines = False
    _x_header(wp, 1, [('Программа', 46, True), ('Буква', 10, True), ('Тип этапа', 16, True),
                      ('Дисциплина', 48, True), ('Часы', 8, True), ('Ставка, руб./час', 16, False),
                      ('Стоимость, руб. (авто)', 20, False)])
    rows = [
        (p1, 'ЦТЭ', 'Занятия', 'Модуль 1: Основы анализа данных', 24, 900),
        (p1, 'ЦТЭ', 'Занятия', 'Модуль 2: Машинное обучение для бизнеса', 16, 900),
        (p1, 'ЦТЭ', 'Сопровождение', None, 40, 499),
        (p2, 'ОПП', 'Занятия', '1. Введение в программирование на Python', 32, 900),
        (p2, 'ОПП', 'Сопровождение', None, 40, 499),
    ]
    for r, row in enumerate(rows, 2):
        _x_row(wp, r, row, zebra=(r % 2 == 1))
        c = wp.cell(row=r, column=7, value=f'=E{r}*F{r}')
        c.font = Font(name=_X_FONT, size=10, color='666666')
        c.border = _X_BORDER
        c.number_format = '#,##0'
        c.alignment = Alignment(vertical='center')
    _x_fit(wp)
    wp.freeze_panes = 'A2'
    wp['A1'].comment = Comment('Название одинаково во всех строках программы; подставляется в {direction}.', 'Генератор')
    wp['C1'].comment = Comment('«Занятия» — строка по дисциплине. «Сопровождение» — организационное и методическое '
                               'сопровождение. Порядок строк = порядок этапов в ТЗ.', 'Генератор')
    wp['F1'].comment = Comment('Если пусто: 900 для занятий, 499 для сопровождения.', 'Генератор')

    _x_notes(wb, 'Как заполнять (договор ГПХ)', [
        '1. Лист «Программы» заполняется один раз на программу. Каждая строка — этап: «Занятия» (дисциплина + часы) или «Сопровождение» (только часы). Порядок строк = порядок этапов в Техническом задании.',
        '2. Лист «Слушатели»: одна строка = один слушатель = один договор. Желтые заголовки — обязательные столбцы: ФИО, Группа, Программа, Сумма договора.',
        '3. Несколько групп у слушателя — через запятую. Этапы в ТЗ повторятся для каждой группы.',
        '4. «Сумма договора» берётся из Excel как есть и подставляется в договор и акт числом и прописью.',
        '5. Название в «Слушатели → Программа» должно совпадать с названием на листе «Программы» (регистр не важен).',
        '6. Верхняя строка с цветными группами («Обязательно», «Личные данные и паспорт» …) — только для удобства, на генерацию не влияет.',
        '7. Строки-примеры замените своими данными. Порядок столбцов можно менять, названия заголовков — нет.',
    ])
    return wb


_EXAMPLE_BUILDERS = {
    'fiz': (_build_fiz_example, 'Primer_tablitsy_fiz.xlsx'),
    'gph': (_build_gph_example, 'Primer_tablitsy_GPKh.xlsx'),
}


class ContractExcelTemplate(View):
    """GET ?type=fiz|gph — отдаёт образец Excel для выбранного типа договора."""

    def get(self, request):
        builder = _EXAMPLE_BUILDERS.get(request.GET.get('type', 'fiz'))
        if builder is None:
            return JsonResponse({'error': 'Неизвестный тип договора'}, status=400)
        build, filename = builder
        buf = io.BytesIO()
        build().save(buf)
        response = HttpResponse(
            buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response